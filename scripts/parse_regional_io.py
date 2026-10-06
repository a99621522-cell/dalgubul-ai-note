#!/usr/bin/env python3
"""한국은행 「2020년 지역산업연관표」 통계표 엑셀(지역간 투입산출표, 생산자가격, 통합대분류)에서 대구 33부문 유발계수를 뽑아
scripts/data/bok_io_daegu_sectors.csv 의 k_va_within(지역내 부가가치유발계수)·k_va_other(타지역)·emp_within(지역내 취업유발계수)·k_prod_within 을 채운다.
운영자 지시 2026-10-06 '업종별 유발계수 찾아서 공장 유치에 간접분 더해'. 엑셀은 운영자가 bok.or.kr 에서 받아 Drive DAITDA 에 올린 것(gdown 으로 러너가 받음, 저장소에 커밋하지 않음).

사용: python3 scripts/parse_regional_io.py --file <Drive id 또는 로컬 xlsx> [--probe] [--out scripts/data/bok_io_daegu_sectors.csv]
--probe: 시트 이름·크기·머리 셀·'대구' 셀 위치만 찍는다(표 구조 확인용). 값은 그대로 옮기고 지역내·타지역 합만 더한다(계수표의 열 합).

지역간 유발계수표 구조(가정, probe 로 확인): 행 = (부가가치·취업이 생기는 지역, 부문), 열 = (최종수요가 생긴 지역, 상품 부문).
대구 상품 s 의 최종수요 1단위 → 열 (대구, s) 의 행 합: 행 지역 = 대구 → 지역내 계수, 그 밖 → 타지역 계수.
"""
import argparse, csv, re, subprocess, sys, tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
REGIONS = ["서울", "인천", "경기", "대전", "세종", "충북", "충남", "광주", "전북", "전남", "대구", "경북", "부산", "울산", "경남", "강원", "제주"]
NORM = lambda s: re.sub(r"[\s,·]", "", str(s or ""))


def fetch(item: str) -> Path:
    p = Path(item)
    if p.exists():
        return p
    d = Path(tempfile.mkdtemp())
    out = d / "io.xlsx"
    subprocess.run(["gdown", item, "-O", str(out)], check=True)   # 최신 gdown 은 --id 없이 id 를 바로 받는다
    return out


def region_of(s) -> str | None:
    t = NORM(s)
    for r in REGIONS:
        if t == r or t.startswith(r):
            return r
    return None


def probe(path: Path) -> None:
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    for ws in wb.worksheets:
        print(f"== 시트 '{ws.title}' {ws.max_row}행 × {ws.max_column}열")
        rows = list(ws.iter_rows(min_row=1, max_row=8, max_col=10, values_only=True))
        for i, r in enumerate(rows, 1):
            print("  ", i, [str(v)[:14] if v is not None else "" for v in r])
        hits = []
        for i, r in enumerate(ws.iter_rows(min_row=1, max_row=min(ws.max_row, 700), max_col=min(ws.max_column, 700), values_only=True), 1):
            for j, v in enumerate(r, 1):
                if isinstance(v, str) and "대구" in v:
                    hits.append((i, j, v[:20]))
        print("   '대구' 셀:", hits[:12], "…" if len(hits) > 12 else "", f"총 {len(hits)}")


def grid(ws):
    return [list(r) for r in ws.iter_rows(values_only=True)]


def num(v):
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(",", ""))
    except (TypeError, ValueError):
        return None


def parse_matrix(ws, label: str) -> dict[str, dict[str, float]]:
    """지역간 계수표 하나 → {부문명: {'within': 합, 'other': 합, 'col_total': 합}} (대구 열 기준).
    확인된 구조(2026-10-06 probe): 5행 = 열 지역(서울×33 …), 6행 = 부문 코드, 7행 = '지역'·'부문'·'부문명' + 열 부문명, 8행부터 = 행 지역·코드·부문명·값."""
    g = grid(ws)
    nrow, ncol = len(g), max(len(r) for r in g)
    cell = lambda i, j: g[i][j] if i < nrow and j < len(g[i]) else None
    reg_row = next((i for i in range(min(15, nrow)) if sum(1 for v in g[i] if region_of(v)) >= 5), None)
    if reg_row is None:
        raise RuntimeError(f"{label}: 열 머리의 지역 행을 못 찾음")
    name_row = next((i for i in range(reg_row + 1, min(reg_row + 5, nrow)) if any(NORM(v) == "부문명" for v in g[i] if isinstance(v, str))), reg_row + 1)
    hdr = [NORM(v) if isinstance(v, str) else "" for v in g[name_row]]
    reg_col = hdr.index("지역") if "지역" in hdr else 0
    name_col = hdr.index("부문명") if "부문명" in hdr else reg_col + 2
    col_region, cur = {}, None
    for j in range(ncol):
        r = region_of(cell(reg_row, j))
        if r:
            cur = r
        col_region[j] = cur
    col_sector = {j: (str(cell(name_row, j)).strip() if cell(name_row, j) is not None else "") for j in range(ncol)}
    out: dict[str, dict[str, float]] = {}
    cur = None
    for i in range(name_row + 1, nrow):
        rr = region_of(cell(i, reg_col))
        if rr:
            cur = rr
        sec = str(cell(i, name_col)).strip() if cell(i, name_col) is not None else ""
        if not cur or not sec or NORM(sec) in ("합계", "계", "총계", "전부문", "부문계", "전품목", "None"):
            continue
        for j in range(name_col + 1, ncol):
            if col_region.get(j) != "대구":
                continue
            cs = col_sector[j]
            if not cs or NORM(cs) in ("합계", "계", "총계", "부문계", "전품목"):
                continue
            v = num(cell(i, j))
            if v is None:
                continue
            d = out.setdefault(cs, {"within": 0.0, "other": 0.0, "col_total": 0.0})
            d["col_total"] += v
            d["within" if cur == "대구" else "other"] += v
    return out


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--file", required=True)
    ap.add_argument("--probe", action="store_true")
    ap.add_argument("--out", default=str(ROOT / "scripts/data/bok_io_daegu_sectors.csv"))
    a = ap.parse_args()
    path = fetch(a.file)
    if a.probe:
        probe(path); return 0
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    pick = lambda words: next((ws for ws in wb.worksheets if all(w in NORM(ws.title) for w in words)), None)
    sheets = {"va": pick(["부가가치유발"]), "emp": pick(["취업유발"]), "prod": pick(["생산유발"])}
    print("시트:", {k: (v.title if v else None) for k, v in sheets.items()})
    res = {k: (parse_matrix(v, k) if v else {}) for k, v in sheets.items()}
    imp = parse_matrix(pick(["수입유발"]), "imp") if pick(["수입유발"]) else {}
    for k, d in res.items():
        print(f"[{k}] 대구 열 {len(d)}개 부문", {s: (round(x['within'], 3), round(x['other'], 3)) for s, x in list(d.items())[:6]})
    for s, x in res["va"].items():   # 검산: 부가가치유발(지역내+타지역) + 수입유발(전 지역) ≈ 1
        tot = x["col_total"] + (imp.get(s, {}).get("col_total", 0))
        if abs(tot - 1) > 0.02:
            print(f"   검산 주의 {s}: 부가가치 {x['col_total']:.3f} + 수입 {imp.get(s, {}).get('col_total', 0):.3f} = {tot:.3f}")
    # CSV 갱신(sector 이름을 공백·쉼표 제거로 맞춤)
    rows = list(csv.DictReader(open(a.out, encoding="utf-8")))
    fields = list(rows[0].keys())
    for extra in ("k_prod_within", "k_prod_other"):
        if extra not in fields:
            fields.append(extra)
    key = {NORM(s): s for s in res["va"]}
    filled = 0
    for r in rows:
        k = key.get(NORM(r["sector"]))
        if not k:
            continue
        va = res["va"].get(k); emp = res["emp"].get(k); pr = res["prod"].get(k)
        if va: r["k_va_within"] = f"{va['within']:.4f}"; r["k_va_other"] = f"{va['other']:.4f}"; filled += 1
        if emp: r["emp_within"] = f"{emp['within']:.3f}"
        if pr: r["k_prod_within"] = f"{pr['within']:.4f}"; r["k_prod_other"] = f"{pr['other']:.4f}"
    unmatched = [s for s in res["va"] if NORM(s) not in {NORM(r["sector"]) for r in rows}]
    print(f"채움 {filled}개 부문, 엑셀에만 있는 부문: {unmatched}")
    with open(a.out, "w", encoding="utf-8", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=fields); w.writeheader()
        for r in rows: w.writerow({f: r.get(f, "") for f in fields})
    return 0


if __name__ == "__main__":
    sys.exit(main())
