#!/usr/bin/env python3
"""글(src/content/posts/<id>.md) → 편집용 Word(.docx).

운영자가 한글·워드에서 고친 뒤 다시 올릴 수 있게, 본문의 제목·문단·표·그림·캡션·주석을
그대로 옮긴다. 그림(public/figures/.../*.svg)은 rsvg-convert 로 PNG 로 바꿔 넣고,
그림마다 대체 텍스트에 원래 파일 경로를 적어 되돌릴 때 어느 그림인지 알 수 있게 한다.

  python3 scripts/post_to_docx.py 2026-09-30-policy-structure-daegu [-o out.docx] [--public]   # --public: 게시용(초안 표시·FAQ 없음)

필요: python-docx, markdown, beautifulsoup4, rsvg-convert(librsvg2-bin)
"""
from __future__ import annotations

import argparse
import re
import subprocess
import sys
import tempfile
from pathlib import Path

import markdown
import yaml
from bs4 import BeautifulSoup, NavigableString, Tag
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
FONT = "맑은 고딕"
KNOWN_TAGS = "figure|img|figcaption|sup|sub|br|strong|em|b|i|a|table|thead|tbody|tr|th|td|p|span|div"


def set_font(run, size=None, bold=None, color=None, italic=None):
    run.font.name = FONT
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for k in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(k), FONT)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def style_base(doc):
    for name in ("Normal", "Heading 1", "Heading 2", "Heading 3", "Title", "Caption", "List Bullet", "Quote"):
        try:
            st = doc.styles[name]
        except KeyError:
            continue
        st.font.name = FONT
        rpr = st.element.get_or_add_rPr()
        fonts = rpr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            rpr.insert(0, fonts)
        for k in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            fonts.set(qn(k), FONT)
        # 제목 스타일의 테마 글꼴 지정을 지워야 맑은 고딕이 먹는다
        for k in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            if fonts.get(qn(k)) is not None:
                del fonts.attrib[qn(k)]
    normal = doc.styles["Normal"]
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.3
    sizes = {"Title": 18, "Heading 1": 15, "Heading 2": 13, "Heading 3": 11.5}
    for name, sz in sizes.items():
        st = doc.styles[name]
        st.font.size = Pt(sz)
        st.font.bold = True
        st.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)
        st.paragraph_format.space_before = Pt(14 if name != "Heading 3" else 10)
        st.paragraph_format.space_after = Pt(6)
    cap = doc.styles["Caption"]
    cap.font.size = Pt(9)
    cap.font.italic = False
    cap.font.color.rgb = RGBColor(0x44, 0x44, 0x44)


def add_inline(par, node, size=None, bold=False, sup=False, color=None):
    """HTML 인라인 노드를 run 으로."""
    if isinstance(node, NavigableString):
        text = str(node)
        if not text:
            return
        text = re.sub(r"\s*\n\s*", " ", text)
        r = par.add_run(text)
        set_font(r, size=size, bold=bold or None, color=color)
        if sup:
            r.font.superscript = True
        return
    if not isinstance(node, Tag):
        return
    name = node.name
    if name == "br":
        par.add_run().add_break()
        return
    b = bold or name in ("strong", "b")
    s = sup or name == "sup"
    for ch in node.children:
        add_inline(par, ch, size=size, bold=b, sup=s, color=color)
    if name == "a" and node.get("href", "").startswith("http") and node.get_text(strip=True) != node["href"]:
        r = par.add_run(f" ({node['href']})")
        set_font(r, size=(size or 10.5) - 1, color="555555")


def shade(cell, fill):
    tcpr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def add_table(doc, tbl):
    rows = tbl.find_all("tr")
    if not rows:
        return
    ncol = max(len(r.find_all(["th", "td"])) for r in rows)
    t = doc.add_table(rows=len(rows), cols=ncol)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, tr in enumerate(rows):
        cells = tr.find_all(["th", "td"])
        for j, c in enumerate(cells):
            cell = t.cell(i, j)
            par = cell.paragraphs[0]
            par.paragraph_format.space_after = Pt(0)
            par.paragraph_format.line_spacing = 1.15
            align = (c.get("style") or "") + " " + (c.get("align") or "")
            if "right" in align:
                par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            elif "center" in align or c.name == "th":
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for ch in c.children:
                add_inline(par, ch, size=9, bold=(c.name == "th"))
            if c.name == "th":
                shade(cell, "E8ECF2")
    # 머리 행은 쪽이 넘어가도 반복
    if rows[0].find("th"):
        trpr = t.rows[0]._tr.get_or_add_trPr()
        h = OxmlElement("w:tblHeader")
        h.set(qn("w:val"), "true")
        trpr.append(h)
    doc.add_paragraph()


def svg_to_png(svg: Path, tmp: Path) -> Path | None:
    out = tmp / (svg.stem + ".png")
    try:
        subprocess.run(["rsvg-convert", "-z", "2", "-b", "white", "-o", str(out), str(svg)],
                       check=True, capture_output=True)
        return out
    except (OSError, subprocess.CalledProcessError) as e:
        print(f"  그림 변환 실패 {svg}: {e}", file=sys.stderr)
        return None


def add_figure(doc, fig, tmp):
    img = fig.find("img")
    cap = fig.find("figcaption")
    src = img.get("src", "") if img else ""
    path = ROOT / "public" / src.lstrip("/")
    pic = None
    if path.suffix == ".svg" and path.exists():
        pic = svg_to_png(path, tmp)
    elif path.exists():
        pic = path
    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.paragraph_format.keep_with_next = True
    if pic:
        run = par.add_run()
        run.add_picture(str(pic), width=Cm(16))
        # 대체 텍스트 = 원래 파일 경로 (되돌릴 때 그림을 찾는 표시)
        for dp in run._element.iter(qn("wp:docPr")):
            dp.set("descr", src)
            dp.set("title", img.get("alt", ""))
    else:
        r = par.add_run(f"[그림 파일 없음: {src}]")
        set_font(r, color="AA0000")
    if cap:
        cp = doc.add_paragraph(style="Caption")
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for ch in cap.children:
            add_inline(cp, ch, size=9, color="444444")


def add_block(doc, el, tmp):
    if isinstance(el, NavigableString):
        if str(el).strip():
            p = doc.add_paragraph()
            add_inline(p, el)
        return
    name = el.name
    if name in ("h1", "h2", "h3", "h4"):
        level = {"h1": 1, "h2": 1, "h3": 2, "h4": 3}[name]
        h = doc.add_heading(level=level)
        for ch in el.children:
            add_inline(h, ch)
    elif name == "p":
        if el.find("figure"):
            for ch in el.children:
                add_block(doc, ch, tmp)
            return
        p = doc.add_paragraph()
        txt = el.get_text()
        if re.match(r"^\s*(주:|자료:|\d+\)\s)", txt):
            for ch in el.children:
                add_inline(p, ch, size=9, color="444444")
        else:
            for ch in el.children:
                add_inline(p, ch)
    elif name == "figure":
        add_figure(doc, el, tmp)
    elif name == "table":
        add_table(doc, el)
    elif name in ("ul", "ol"):
        style = "List Bullet" if name == "ul" else "List Number"
        for li in el.find_all("li", recursive=False):
            p = doc.add_paragraph(style=style)
            for ch in li.children:
                if isinstance(ch, Tag) and ch.name in ("ul", "ol"):
                    continue
                if isinstance(ch, Tag) and ch.name == "p":
                    for c2 in ch.children:
                        add_inline(p, c2)
                else:
                    add_inline(p, ch)
    elif name == "blockquote":
        for ch in el.children:
            if isinstance(ch, Tag):
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.6)
                ppr = p._element.get_or_add_pPr()
                bdr = OxmlElement("w:pBdr")
                left = OxmlElement("w:left")
                for k, v in (("val", "single"), ("sz", "18"), ("space", "8"), ("color", "3B6EA5")):
                    left.set(qn(f"w:{k}"), v)
                bdr.append(left)
                ppr.append(bdr)
                for c2 in ch.children:
                    add_inline(p, c2)
    elif name == "hr":
        doc.add_paragraph("―" * 20).alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        for ch in el.children:
            add_block(doc, ch, tmp)


def split_footnotes(body: str) -> tuple[str, list[tuple[str, str]]]:
    """각주([^n] 참조와 '[^n]: 설명' 정의)를 정의 순서대로 번호를 매겨 위첨자 [n] 과 '주석' 목록으로 바꾼다.
    숫자 각주는 그 번호(=출처 목록 번호)를 그대로 쓰고, 이름 각주([^corp] 등)는 뒤 번호를 이어 붙인다."""
    defs = re.findall(r"^\[\^([^\]\s]+)\]:[ \t]*(.*)$", body, re.M)
    if not defs:
        return body, []
    nums = [int(k) for k, _ in defs if k.isdigit()]
    nxt = max(nums, default=0) + 1
    label: dict[str, str] = {}
    for k, _ in defs:
        if k.isdigit():
            label[k] = k
        else:
            label[k] = str(nxt)
            nxt += 1
    body = re.sub(r"^\[\^([^\]\s]+)\]:[ \t]*.*\n?", "", body, flags=re.M)
    body = re.sub(r"\[\^([^\]\s]+)\]", lambda m: f"<sup>[{label.get(m.group(1), m.group(1))}]</sup>", body)
    return body, [(label[k], v) for k, v in defs]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("post_id")
    ap.add_argument("-o", "--out")
    ap.add_argument("--public", action="store_true", help="게시용: 초안 표시·파일 id 줄과 FAQ 를 빼고 날짜만 적는다")
    args = ap.parse_args()
    src = ROOT / "src/content/posts" / f"{args.post_id}.md"
    raw = src.read_text(encoding="utf-8")
    m = re.match(r"^---\n(.*?)\n---\n(.*)$", raw, re.S)
    fm, body = yaml.safe_load(m.group(1)), m.group(2)
    # 본문의 '<참고 1>' 같은 꺾쇠 글자를 태그로 읽지 않게
    body = re.sub(rf"<(?!/?(?:{KNOWN_TAGS})\b)", "&lt;", body)
    body, notes = split_footnotes(body)
    html = markdown.markdown(body, extensions=["tables", "md_in_html"])
    soup = BeautifulSoup(html, "html.parser")

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Cm(2.2))
    sec.top_margin, sec.bottom_margin = Cm(2), Cm(2)
    style_base(doc)

    t = doc.add_paragraph(style="Title")
    set_font(t.add_run(fm.get("title", args.post_id)), size=18, bold=True)
    meta = doc.add_paragraph()
    meta_text = str(fm.get('date', '')) if args.public else f"{fm.get('date', '')} · 초안(편집용) · 파일 id {args.post_id}"
    set_font(meta.add_run(meta_text), size=9, color="666666")
    if fm.get("description"):
        p = doc.add_paragraph()
        set_font(p.add_run("핵심 문장  "), size=10, bold=True)
        set_font(p.add_run(str(fm["description"])), size=10)

    with tempfile.TemporaryDirectory() as td:
        tmp = Path(td)
        for el in soup.children:
            add_block(doc, el, tmp)

        if notes:
            doc.add_heading("주석", level=1)
            for n, v in notes:
                p = doc.add_paragraph()
                set_font(p.add_run(f"[{n}] {v}"), size=9.5)
        if fm.get("faq") and not args.public:
            doc.add_heading("자주 묻는 질문", level=1)
            for qa in fm["faq"]:
                p = doc.add_paragraph()
                set_font(p.add_run("Q. " + qa["q"]), bold=True)
                p = doc.add_paragraph()
                set_font(p.add_run("A. " + qa["a"]))
        if fm.get("sources"):
            doc.add_heading("출처", level=1)
            for s in fm["sources"]:
                p = doc.add_paragraph(style="List Number")
                set_font(p.add_run(f"{s.get('title', '')} — {s.get('url', '')} ({s.get('date', '')})"), size=9.5)

        out = Path(args.out) if args.out else ROOT / "docs/drafts" / f"{args.post_id}.docx"
        out.parent.mkdir(parents=True, exist_ok=True)
        doc.save(out)
    print(f"저장: {out} ({out.stat().st_size // 1024} KB)")


if __name__ == "__main__":
    main()
