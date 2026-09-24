#!/usr/bin/env python3
"""팩토리온 내려받기 파일(.xls/.xlsx) 진단·정리

팩토리온 자료실에서 받은 파일이 엑셀에서 안 열리거나 정렬이 안 될 때 쓴다.
확장자가 .xls 여도 실제로는 HTML 표·CSV·XML(SpreadsheetML)인 경우가 많다.

사용:
  python3 scripts/repair_factoryon.py <파일>            # 진단 + 정리본 저장
  python3 scripts/repair_factoryon.py <파일> --inspect  # 진단만 (첫 바이트·실제 형식)
  python3 scripts/repair_factoryon.py                   # data/, scripts/data/ 의 .xls/.xlsx 전부

하는 일:
  1. 첫 바이트로 실제 형식 판정 (OLE2 xls / OOXML xlsx / HTML 표 / SpreadsheetML XML / CSV·TSV 텍스트)
  2. 형식에 맞게 pandas 로 읽고 헤더 행을 찾는다 (제목 행이 위에 있어도 됨)
  3. 원본은 그대로 두고 옆에 `<원본이름>_정리.xlsx` 저장
     - 시트 1개, 첫 행 헤더 고정, 자동 필터, 열 너비 자동
     - '종사자' 열은 숫자, '등록일·일자' 열은 날짜 (못 읽은 값은 원문 그대로 둠)
     - 사업자번호·우편번호 등 나머지는 문자열 유지 (앞자리 0 보존)

필요: pip install pandas openpyxl xlrd beautifulsoup4 --break-system-packages
"""
from __future__ import annotations

import csv
import io
import re
import sys
import xml.etree.ElementTree as ET
from datetime import datetime
from pathlib import Path

import pandas as pd

ROOT = Path(__file__).resolve().parent.parent
SCAN_DIRS = [ROOT / "data", ROOT / "scripts" / "data"]
SUFFIX = "_정리"

# 헤더 행을 찾을 때 쓰는 팩토리온 열 이름 후보
HEADER_HINTS = ("회사명", "공장명", "업체명", "기업명", "업종", "종사자", "소재지", "주소", "단지")
NUMERIC_COLS = re.compile(r"종사자|종업원")
DATE_COLS = re.compile(r"등록일|일자|설립일|계약일|입주일|승인일")
TEXT_ENCODINGS = ("utf-8-sig", "utf-16", "cp949", "euc-kr", "utf-8")


# ---------------------------------------------------------------- 형식 판정
def sniff(path: Path) -> tuple[str, bytes]:
    """실제 형식을 첫 바이트로 판정. (형식, 첫 16바이트)"""
    with open(path, "rb") as f:
        head = f.read(4096)
    magic = head[:16]

    if head.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"):
        return "ole2", magic  # 진짜 BIFF xls (Compound File Binary)
    if head.startswith(b"PK\x03\x04"):
        return "zip", magic  # xlsx (OOXML). 손상 여부는 열어봐야 안다
    if head.startswith(b"\x1f\x8b"):
        return "gzip", magic

    text = _decode_head(head)
    low = text.lstrip("﻿ \t\r\n").lower()
    if low.startswith("<?xml") and "urn:schemas-microsoft-com:office:spreadsheet" in low:
        return "spreadsheetml", magic
    if low.startswith(("<!doctype html", "<html", "<table", "<meta", "<head", "<body")) or "<table" in low:
        return "html", magic
    if low.startswith("<?xml") or low.startswith("<"):
        return "xml", magic

    # 텍스트: 첫 몇 줄에 구분자가 있으면 CSV/TSV
    sample = "\n".join(text.splitlines()[:5])
    if sample.count("\t") >= 2:
        return "tsv", magic
    if sample.count(",") >= 2:
        return "csv", magic
    # 출력 가능한 문자만이면 텍스트로 본다
    printable = sum(ch.isprintable() or ch in "\r\n\t" for ch in text)
    if text and printable / len(text) > 0.95:
        return "text", magic
    return "unknown", magic


def _decode_head(head: bytes) -> str:
    if head.startswith((b"\xff\xfe", b"\xfe\xff")):
        try:
            return head.decode("utf-16", errors="ignore")
        except Exception:
            pass
    for enc in ("utf-8", "cp949"):
        try:
            return head.decode(enc)
        except UnicodeDecodeError:
            continue
    return head.decode("latin-1", errors="ignore")


def read_text(path: Path) -> tuple[str, str]:
    raw = path.read_bytes()
    for enc in TEXT_ENCODINGS:
        try:
            s = raw.decode(enc)
        except UnicodeDecodeError:
            continue
        if enc == "utf-16" and not raw.startswith((b"\xff\xfe", b"\xfe\xff")):
            continue
        return s, enc
    return raw.decode("cp949", errors="replace"), "cp949(대체)"


def describe(kind: str, magic: bytes) -> str:
    names = {
        "ole2": "진짜 xls (OLE2 Compound File, BIFF)",
        "zip": "zip 컨테이너 — xlsx(OOXML)일 가능성",
        "html": "HTML 표 (확장자만 xls)",
        "spreadsheetml": "Excel 2003 XML Spreadsheet (SpreadsheetML)",
        "xml": "XML (스프레드시트 아님)",
        "csv": "CSV 텍스트",
        "tsv": "탭 구분 텍스트",
        "text": "일반 텍스트",
        "gzip": "gzip 압축",
        "unknown": "알 수 없음",
    }
    hexs = " ".join(f"{b:02X}" for b in magic)
    return f"{names.get(kind, kind)}  | 첫 16바이트: {hexs}"


# ---------------------------------------------------------------- 읽기
def load_frame(path: Path, kind: str) -> tuple[pd.DataFrame, str]:
    """형식별로 읽어 (헤더 없는 원시 표, 메모) 반환. 모든 셀은 문자열."""
    if kind == "ole2":
        df = pd.read_excel(path, header=None, dtype=str, engine="xlrd")
        return df, "xlrd"
    if kind == "zip":
        df = pd.read_excel(path, header=None, dtype=str, engine="openpyxl")
        return df, "openpyxl"
    if kind == "html":
        text, enc = read_text(path)
        return _read_html_table(text), f"bs4, 인코딩 {enc}"
    if kind == "spreadsheetml":
        text, enc = read_text(path)
        return _read_spreadsheetml(text), f"SpreadsheetML, 인코딩 {enc}"
    if kind in ("csv", "tsv", "text"):
        text, enc = read_text(path)
        sample = "\n".join(text.splitlines()[:20])
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",\t;|")
            sep = dialect.delimiter
        except csv.Error:
            sep = "\t" if kind == "tsv" else ","
        df = pd.read_csv(io.StringIO(text), sep=sep, header=None, dtype=str, engine="python",
                         quoting=csv.QUOTE_MINIMAL, skip_blank_lines=True)
        return df, f"csv sep={sep!r}, 인코딩 {enc}"
    raise ValueError(f"읽을 수 없는 형식: {kind}")


def _read_html_table(text: str) -> pd.DataFrame:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(text, "html.parser")
    tables = soup.find_all("table")
    if not tables:
        raise ValueError("HTML 안에 <table> 이 없음")
    # 행이 가장 많은 표를 고른다 (제목·안내 표 제외)
    table = max(tables, key=lambda t: len(t.find_all("tr")))
    rows: list[list[str]] = []
    pending: dict[int, tuple[int, str]] = {}  # rowspan 이월: 열 → (남은 행, 값)
    for tr in table.find_all("tr"):
        row: list[str] = []
        col = 0
        cells = tr.find_all(["td", "th"], recursive=False) or tr.find_all(["td", "th"])
        for cell in cells:
            while col in pending:
                left, val = pending[col]
                row.append(val)
                if left <= 1:
                    del pending[col]
                else:
                    pending[col] = (left - 1, val)
                col += 1
            val = " ".join(cell.get_text(" ", strip=True).split())
            try:
                cs = max(1, int(cell.get("colspan", 1)))
            except ValueError:
                cs = 1
            try:
                rs = max(1, int(cell.get("rowspan", 1)))
            except ValueError:
                rs = 1
            for i in range(cs):
                row.append(val if i == 0 else "")
                if rs > 1:
                    pending[col] = (rs - 1, val)
                col += 1
        while col in pending:
            left, val = pending.pop(col)
            row.append(val)
            if left > 1:
                pending[col] = (left - 1, val)
            col += 1
        rows.append(row)
    width = max((len(r) for r in rows), default=0)
    rows = [r + [""] * (width - len(r)) for r in rows]
    return pd.DataFrame(rows, dtype=str)


def _read_spreadsheetml(text: str) -> pd.DataFrame:
    ns = {"ss": "urn:schemas-microsoft-com:office:spreadsheet"}
    root = ET.fromstring(text.encode("utf-8"))
    ws = root.find("ss:Worksheet", ns)
    if ws is None:
        raise ValueError("SpreadsheetML 에 Worksheet 가 없음")
    rows: list[list[str]] = []
    for r in ws.iter(f"{{{ns['ss']}}}Row"):
        row: list[str] = []
        for c in r.findall("ss:Cell", ns):
            idx = c.get(f"{{{ns['ss']}}}Index")
            if idx:
                row.extend([""] * (int(idx) - 1 - len(row)))
            d = c.find("ss:Data", ns)
            row.append("".join(d.itertext()).strip() if d is not None else "")
        rows.append(row)
    width = max((len(r) for r in rows), default=0)
    rows = [r + [""] * (width - len(r)) for r in rows]
    return pd.DataFrame(rows, dtype=str)


# ---------------------------------------------------------------- 정리
def find_header_row(raw: pd.DataFrame) -> int:
    """제목 행·빈 행을 건너뛰고 열 이름이 있는 행을 찾는다."""
    best, best_score = 0, -1
    for i in range(min(len(raw), 30)):
        cells = [str(v).strip() for v in raw.iloc[i].tolist()]
        filled = [c for c in cells if c and c.lower() != "nan"]
        if len(filled) < 2:
            continue
        hint = sum(any(h in c for h in HEADER_HINTS) for c in filled)
        score = hint * 10 + len(filled)
        if hint and score > best_score:
            best, best_score = i, score
    return best


def tidy(raw: pd.DataFrame) -> tuple[pd.DataFrame, dict]:
    raw = raw.fillna("").astype(str)
    raw = raw.replace({"nan": "", "None": ""})
    h = find_header_row(raw)
    names = [c.strip() for c in raw.iloc[h].tolist()]
    seen: dict[str, int] = {}
    cols: list[str] = []
    for i, n in enumerate(names):
        n = n or f"열{i + 1}"
        if n in seen:
            seen[n] += 1
            n = f"{n}_{seen[n]}"
        else:
            seen[n] = 1
        cols.append(n)
    df = raw.iloc[h + 1:].copy()
    df.columns = cols
    df = df.apply(lambda s: s.str.strip())
    df = df[~(df == "").all(axis=1)]  # 빈 행 제거
    # 완전히 빈 열(이름도 없고 값도 없음) 제거
    empty = [c for c in df.columns if c.startswith("열") and (df[c] == "").all()]
    df = df.drop(columns=empty).reset_index(drop=True)

    info = {"header_row": h, "numeric": {}, "date": {}}
    for c in df.columns:
        if NUMERIC_COLS.search(c):
            df[c], bad = to_number(df[c])
            info["numeric"][c] = bad
        elif DATE_COLS.search(c):
            df[c], bad = to_date(df[c])
            info["date"][c] = bad
    return df, info


def to_number(s: pd.Series) -> tuple[pd.Series, int]:
    cleaned = s.str.replace(r"[,\s명인]", "", regex=True)
    num = pd.to_numeric(cleaned, errors="coerce")
    whole = num.dropna()
    if len(whole) and (whole == whole.round()).all():
        num = num.astype("Int64").astype(object)
    out = num.where(num.notna(), s.where(s != "", None))
    bad = int(((num.isna()) & (s != "")).sum())
    return out, bad


def to_date(s: pd.Series) -> tuple[pd.Series, int]:
    def parse(v: str):
        v = v.strip()
        if not v:
            return None
        digits = re.sub(r"\D", "", v)
        if re.fullmatch(r"\d{8}", digits) and re.fullmatch(r"\d{4}[.\-/년 ]*\d{1,2}[.\-/월 ]*\d{1,2}[일 ]*", v):
            try:
                return datetime.strptime(digits, "%Y%m%d")
            except ValueError:
                return None
        m = re.fullmatch(r"(\d{4})[.\-/년 ]+(\d{1,2})[.\-/월 ]+(\d{1,2})[일 ]*", v)
        if m:
            try:
                return datetime(int(m[1]), int(m[2]), int(m[3]))
            except ValueError:
                return None
        if re.fullmatch(r"\d{4,6}(\.0)?", v):  # 엑셀 일련번호 (xls 에서 숫자로 읽힌 경우)
            try:
                return (pd.Timestamp("1899-12-30") + pd.Timedelta(days=int(float(v)))).to_pydatetime()
            except Exception:
                return None
        t = pd.to_datetime(v, errors="coerce")
        return None if pd.isna(t) else t.to_pydatetime()

    # Series.map 은 None 을 NaT 로 바꿔 버리므로 파이썬 리스트로 다룬다
    parsed = [parse(v) for v in s.tolist()]
    out = pd.Series([p if p is not None else (v if v else None) for p, v in zip(parsed, s)], index=s.index, dtype=object)
    bad = sum(1 for p, v in zip(parsed, s) if p is None and v)
    return out, bad


# ---------------------------------------------------------------- 저장
def save_xlsx(df: pd.DataFrame, out: Path, info: dict) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "입주업체"
    font = Font(name="맑은 고딕", size=10)
    head_font = Font(name="맑은 고딕", size=10, bold=True)
    head_fill = PatternFill("solid", fgColor="DDEBF7")

    ws.append(list(df.columns))
    for cell in ws[1]:
        cell.font = head_font
        cell.fill = head_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    numeric_cols = set(info["numeric"])
    date_cols = set(info["date"])
    for row in df.itertuples(index=False):
        ws.append([None if (v is None or (isinstance(v, float) and pd.isna(v))) else v for v in row])

    for j, name in enumerate(df.columns, start=1):
        col = get_column_letter(j)
        fmt = "#,##0" if name in numeric_cols else ("yyyy-mm-dd" if name in date_cols else "@")
        for cell in ws[col][1:]:
            cell.font = font
            if isinstance(cell.value, (int, float)) and name in numeric_cols:
                cell.number_format = fmt
            elif isinstance(cell.value, datetime):
                cell.number_format = "yyyy-mm-dd"
            else:
                cell.number_format = "@"
        # 열 너비: 한글은 2칸으로 세고, 6~60 사이
        lengths = [_display_len(name)]
        for v in df[name].head(2000):
            if isinstance(v, datetime):
                lengths.append(10)
            elif v is not None and not (isinstance(v, float) and pd.isna(v)):
                lengths.append(_display_len(str(v)))
        ws.column_dimensions[col].width = max(6, min(60, max(lengths) + 2))

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(out)


def _display_len(s: str) -> int:
    return sum(2 if ord(ch) > 0x2E7F else 1 for ch in s)


# ---------------------------------------------------------------- 실행
def process(path: Path, inspect_only: bool = False) -> Path | None:
    kind, magic = sniff(path)
    print(f"[{path.name}] {describe(kind, magic)}")
    if inspect_only:
        return None
    try:
        raw, how = load_frame(path, kind)
    except Exception as e:  # 형식 판정이 맞아도 파일이 깨졌을 수 있다
        print(f"  읽기 실패 ({how_hint(kind)}): {e}")
        return None
    df, info = tidy(raw)
    out = path.with_name(f"{path.stem}{SUFFIX}.xlsx")
    save_xlsx(df, out, info)
    print(f"  읽기: {how} | 헤더 {info['header_row'] + 1}행 | {len(df):,}행 × {len(df.columns)}열")
    for c, bad in info["numeric"].items():
        print(f"  숫자 열 '{c}'" + (f" — 숫자로 못 읽은 값 {bad}건은 원문 유지" if bad else ""))
    for c, bad in info["date"].items():
        print(f"  날짜 열 '{c}'" + (f" — 날짜로 못 읽은 값 {bad}건은 원문 유지" if bad else ""))
    print(f"  저장: {out}")
    return out


def how_hint(kind: str) -> str:
    return {"ole2": "xlrd", "zip": "openpyxl", "html": "bs4"}.get(kind, kind)


def main(argv: list[str]) -> int:
    inspect_only = "--inspect" in argv
    args = [a for a in argv if not a.startswith("--")]
    if args:
        paths = [Path(a) for a in args]
    else:
        paths = [p for d in SCAN_DIRS if d.is_dir()
                 for p in sorted(d.iterdir())
                 if p.suffix.lower() in (".xls", ".xlsx") and not p.stem.endswith(SUFFIX)]
        if not paths:
            print("정리할 .xls/.xlsx 가 없음: " + ", ".join(str(d) for d in SCAN_DIRS))
            return 1
    rc = 0
    for p in paths:
        if not p.is_file():
            print(f"[{p}] 파일 없음")
            rc = 1
            continue
        if process(p, inspect_only) is None and not inspect_only:
            rc = 1
    return rc


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
